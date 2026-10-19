"""Textile three-dimension risk POC; source facts and trial policy are separate.

Client = Return Qty/IV product Qty. Product = Strategy. Supplier = PQC/AQL.
Unknown weighted inputs stay unknown: show the known contribution and a bound,
never silently allocate their weight to another input. Scores are priorities,
not defect probabilities. The defaults are proposals, not calibrated standards.
"""
from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path
import math
import re

import numpy as np
import pandas as pd

from dashboard_scope import DashboardScope, filter_records, identifiers, customer_totals

RISK_POC_VERSION = "2026-10-19-v5-return-quantity-1.8-mean"
DEPARTMENTS = ("Hugss", "Knit", "Woven")
GRAINS = ("Model", "CC", "Supplier")


@dataclass(frozen=True)
class RiskProfile:
    client_weight: float = 40
    product_weight: float = 10
    supplier_weight: float = 50
    return_weight: float = 50
    iv_weight: float = 50
    pqc_weight: float = 80
    aql_weight: float = 20
    pqc_warning_pct: float = 4
    pqc_critical_pct: float = 8
    pqc_prior_defects: float = 8
    pqc_prior_samples: float = 200
    aql_warning_pct: float = 10
    aql_critical_pct: float = 20
    return_critical_multiple: float = 1.8
    iv_warning: float = 5
    iv_critical: float = 10

    def validate(self):
        for weights in ((self.client_weight, self.product_weight, self.supplier_weight),
                        (self.return_weight, self.iv_weight), (self.pqc_weight, self.aql_weight)):
            if any(not math.isfinite(v) or v < 0 for v in weights) or not math.isclose(sum(weights), 100):
                raise ValueError("Each weight group must contain non-negative values summing to 100%")
        for warning, critical in ((self.pqc_warning_pct, self.pqc_critical_pct),
                                  (self.aql_warning_pct, self.aql_critical_pct),
                                  (self.iv_warning, self.iv_critical)):
            if not (math.isfinite(warning) and math.isfinite(critical) and 0 < warning < critical):
                raise ValueError("Every warning threshold must be positive and below its critical threshold")
        if not math.isfinite(self.return_critical_multiple) or self.return_critical_multiple <= 1:
            raise ValueError("Return Qty critical multiple must exceed one")
        if not (math.isfinite(self.pqc_prior_samples) and math.isfinite(self.pqc_prior_defects)
                and self.pqc_prior_samples >= 0 and self.pqc_prior_defects >= 0):
            raise ValueError("PQC smoothing inputs must be finite and non-negative")
        return self


def default_profile(grain="Model"):
    if grain not in GRAINS:
        raise ValueError("Unsupported risk grain")
    # Strategy is a product characteristic, not a whole-supplier attribute.
    return replace(RiskProfile(), product_weight=0, supplier_weight=60) if grain == "Supplier" else RiskProfile()


def profile_from_saved(saved, grain="Model"):
    """Migrate session policies without reusing RPM thresholds as quantities."""
    policy = dict(saved or {})
    if "return_weight" not in policy and "rpm_weight" in policy:
        policy["return_weight"] = policy["rpm_weight"]
    return replace(default_profile(grain), **{
        name: value for name, value in policy.items() if name in RiskProfile.__dataclass_fields__
    })


def number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def threshold_score(value, warning, critical):
    value = number(value)
    if value is None or value < 0:
        return None
    if not 0 < warning < critical:
        raise ValueError("Invalid scoring thresholds")
    return min(100., 50 * value / warning if value <= warning else 50 + 50 * (value-warning)/(critical-warning))


def weighted_bounds(parts):
    """Combine (score/None, percent weight); zero-weight missing parts are valid."""
    known = coverage = 0.
    for score, weight in parts:
        if weight == 0:
            continue
        score = number(score)
        if score is not None:
            known += min(100., max(0., score)) * weight / 100
            coverage += weight
    missing = max(0., 100.-coverage)
    return {"score": known if math.isclose(missing, 0) else None,
            "lower": known, "upper": min(100., known+missing), "coverage": coverage}


def customer_return_qty(frame):
    """Return quantities are independent of whether a sales denominator exists."""
    if frame.empty or "returned_now" not in frame:
        return None
    grains = set(frame.get("customer_grain", pd.Series("", index=frame.index)))
    if len(grains) != 1 or next(iter(grains)) not in ("Model", "CC"):
        return None
    key = "model" if next(iter(grains)) == "Model" else "cc"
    if key not in frame or frame[key].eq("").any():
        return None
    values = frame.copy()
    values["_qty"] = pd.to_numeric(values.returned_now, errors="coerce")
    if not (values._qty.ge(0) & np.isfinite(values._qty)).all():
        return None
    unique = values.groupby(["supplier_code", key])._qty.agg(lambda group: group.iloc[0] if group.nunique() == 1 else np.nan)
    return float(unique.sum()) if unique.notna().all() else None


def iv_quantity_facts(frame, *, available=True):
    """Deduplicate a feedback/product quantity; never substitute its row count."""
    if not available or "iv_qty" not in frame:
        return {"iv": None, "iv_records": None}
    values = frame.copy()
    values["_id"] = identifiers(values.get("case_id", pd.Series("", index=values.index)))
    values["_qty"] = pd.to_numeric(values.iv_qty, errors="coerce")
    if values._id.eq("").any() or not (values._qty.ge(0) & np.isfinite(values._qty)).all():
        return {"iv": None, "iv_records": int(values._id[values._id.ne("")].nunique())}
    values["_model"] = identifiers(values.get("model", values.get("model_code", pd.Series("", index=values.index))))
    quantities = values.groupby(["_id", "_model"])._qty.agg(lambda group: group.iloc[0] if group.nunique() == 1 else np.nan)
    return {"iv": float(quantities.sum()) if quantities.notna().all() else None, "iv_records": int(values._id.nunique())}


def return_qty_reference(customer, population, key):
    """Same-grain peer mean includes explicit zero and excludes unknown returns."""
    peer_values = {}
    for entity in sorted(population):
        value = customer_return_qty(customer.loc[customer[key].eq(entity)])
        if value is not None:
            peer_values[entity] = value
    total = sum(peer_values.values())
    return {"mean": total/len(peer_values) if peer_values else None,
            "count": len(peer_values), "total": total, "values": peer_values,
            "population": len(population), "key": key}


def return_quantity_score(value, reference, profile):
    mean = number(reference.get("mean"))
    value = number(value)
    if value is None or mean is None or value < 0:
        return None
    if mean == 0:
        return 0. if value == 0 else None
    return threshold_score(value, mean, mean*profile.return_critical_multiple)


def identifier(value, digits):
    text = "" if value is None or pd.isna(value) else str(value).strip()
    text = re.sub(r"\.0$", "", text)
    return text if re.fullmatch(r"\d{"+str(digits)+r"}", text) else ""


def id_tokens(value, digits):
    """Only a numeric ID or an explicit separator-delimited ID list is usable."""
    if value is None or pd.isna(value):
        return frozenset()
    pieces = re.split(r"[,;/|\s、，；]+", str(value).strip())
    codes = [identifier(piece, digits) for piece in pieces if piece]
    return frozenset(codes) if codes and all(codes) else frozenset()


def load_strategy_workbook(path):
    """Read only the formal follow-up input sheet, not archives or formula tabs.

    The supplied follow-up is partial; the user's current POC policy treats unlisted Models as 0.
    PE move, planned milestones, sport and development codes do not add risk.
    Preserve row-level ID sets without inventing CC/Model Cartesian links.
    """
    from openpyxl import load_workbook
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook["Strategic FG homologation"]
        header = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
        if not all(term in str(header[index]).casefold() for index, term in
                   ((8, "supplier"), (11, "conception"), (12, "model"), (14, "season"))):
            raise ValueError("Strategic workbook schema changed")
        records = []
        for row_num, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            models, ccs = id_tokens(row[12], 7), id_tokens(row[11], 6)
            if not models and not ccs:
                continue
            supplier = re.match(r"^\s*(\d{5,6})(?:\s*-|\s*$)", str(row[8] or ""))
            records.append({"models": sorted(models), "ccs": sorted(ccs),
                            "supplier_code": supplier.group(1) if supplier else "",
                            "season": str(row[14] or "").strip().upper(),
                            "strategic": True, "source_row": row_num,
                            "source": "Strategic FG homologation"})
        return records
    finally:
        workbook.close()


def strategic_identity(models, cc, registry, supplier_code="49425", season=""):
    """Binary list membership: any exact strategic Model makes its CC strategic.

    Supplier identity is mandatory. An explicitly selected season narrows the
    list; no selected season means all available seasons. Unlisted = 0 under
    the user's POC policy. This flag does not modify the two chart coordinates.
    """
    models = set(models) - {""}
    matches = [record for record in registry
               if record.get("supplier_code") == supplier_code
               and (not season or str(record.get("season", "")).upper() == season.upper())
               and bool(models.intersection(record.get("models", ())))]
    strategic = any(record.get("strategic") is True for record in matches)
    return {"status": "yes" if strategic else "no", "indicator": int(strategic),
            "score": 100. if strategic else 0.,
            "matched_models": len({model for record in matches for model in record.get("models", ()) if model in models}),
            "model_count": len(models), "reason": "Binary Strategy list membership; any listed strategic Model applies to the CC",
            "source_rows": sorted({record["source_row"] for record in matches if record.get("source_row") is not None})}


def normalize_source(frame, kind):
    result = frame.copy()
    cc = result.get("product_code", result.get("cc", pd.Series("", index=result.index)))
    model = result.get("model_code", result.get("model", pd.Series("", index=result.index)))
    if kind == "pqc":
        fallback = result.get("product_label", pd.Series("", index=result.index)).fillna("").astype(str).str.extract(r"^(\d{7})(?!\d)", expand=False)
        model = model.where(model.fillna("").astype(str).str.strip().ne(""), fallback)
    result["cc"] = cc.map(lambda value: identifier(value, 6))
    result["model"] = model.map(lambda value: identifier(value, 7))
    # This adapter is deliberately ZX/Hugss-only. More factories require an
    # authoritative department and supplier mapping, not process-name guesses.
    result["supplier_code"] = "49425"
    return result


def pqc_facts(frame, profile):
    d = pd.to_numeric(frame.get("defect_qty", pd.Series(dtype=float)), errors="coerce")
    n = pd.to_numeric(frame.get("qty_inspected", pd.Series(dtype=float)), errors="coerce")
    valid = d.ge(0) & n.gt(0) & np.isfinite(d) & np.isfinite(n)
    complete = len(frame) > 0 and bool(valid.all())
    defects, inspected = float(d[valid].sum()), float(n[valid].sum())
    smoothed = ((defects + profile.pqc_prior_defects) / (inspected + profile.pqc_prior_samples)) if complete else None
    return {"pqc_defects": defects, "pqc_inspected": inspected, "pqc_rows": len(frame),
            "pqc_valid_rows": int(valid.sum()), "pqc_dpu": defects/inspected if complete else None,
            "pqc_smoothed": smoothed, "pqc_score": threshold_score(smoothed*100, profile.pqc_warning_pct, profile.pqc_critical_pct) if smoothed is not None else None}


def aql_facts(frame, profile):
    frame = frame.copy()
    # Ignore exact duplicate export records; conflicting copies of a record ID
    # make its first-result status unresolved, not another independent check.
    if "record_id" in frame:
        frame["_id"] = identifiers(frame.record_id)
        frame.loc[frame._id.eq(""), "_id"] = [f"missing-{i}" for i in frame.index[frame._id.eq("")]]
    else:
        frame["_id"] = [f"row-{i}" for i in range(len(frame))]
    text = frame.get("result_raw", frame.get("result", pd.Series("", index=frame.index))).fillna("").astype(str)
    failed = text.str.contains(r"FAIL|NG|NOK|不合格|拒|RE-CHECK|RECHECK|重验", case=False, regex=True)
    passed = text.str.contains(r"PASS|合格", case=False, regex=True) & ~failed
    frame["_status"] = np.where(failed, "fail", np.where(passed, "pass", "unknown"))
    states = frame.groupby("_id")["_status"].agg(lambda values: next(iter(set(values))) if len(set(values)) == 1 else "unknown")
    total, valid, fail = len(states), int(states.ne("unknown").sum()), int(states.eq("fail").sum())
    rate = fail/valid if total and valid == total else None
    return {"aql_records": total, "aql_valid_records": valid, "aql_fail": fail,
            "aql_fail_rate": rate, "aql_score": threshold_score(rate*100, profile.aql_warning_pct, profile.aql_critical_pct) if rate is not None else None}


def score_entity(pqc, aql, customer, cases, profile, strategy, *, iv_available=True, rpm_available=True, return_reference=None):
    profile.validate()
    facts = {**pqc_facts(pqc, profile), **aql_facts(aql, profile)}
    # Exact duplicates are harmless. Conflicting duplicate rows at one selected
    # customer grain are not additional sales and cannot be averaged.
    cust = customer.drop_duplicates()
    duplicate = bool(cust.duplicated(["supplier_code", "model" if cust.get("customer_grain", pd.Series(dtype=str)).eq("Model").all() else "cc"]).any()) if not cust.empty else False
    totals = customer_totals(cust) if rpm_available and not duplicate else {}
    rpm = totals.get("rpm_now")
    returned = customer_return_qty(customer) if rpm_available else None
    iv_facts = iv_quantity_facts(cases, available=iv_available)
    iv = iv_facts["iv"]
    reference = return_reference or {}
    return_score = return_quantity_score(returned, reference, profile)
    iv_score = threshold_score(iv, profile.iv_warning, profile.iv_critical)
    client = weighted_bounds(((return_score, profile.return_weight), (iv_score, profile.iv_weight)))
    supplier = weighted_bounds(((facts["pqc_score"], profile.pqc_weight), (facts["aql_score"], profile.aql_weight)))
    product_score = strategy["score"]
    # Propagate uncertainty through nested weights, rather than discarding a
    # partly measured dimension's known contribution.
    lower = client["lower"]*profile.client_weight/100 + (product_score or 0)*profile.product_weight/100 + supplier["lower"]*profile.supplier_weight/100
    upper = client["upper"]*profile.client_weight/100 + (100 if product_score is None else product_score)*profile.product_weight/100 + supplier["upper"]*profile.supplier_weight/100
    coverage = client["coverage"]*profile.client_weight/100 + (0 if product_score is None else profile.product_weight) + supplier["coverage"]*profile.supplier_weight/100
    facts.update({"rpm": rpm, "returns": returned, "sales": totals.get("sold_now"), **iv_facts,
                  "return_qty_score": return_score, "return_qty_mean": reference.get("mean"),
                  "return_reference_count": reference.get("count", 0), "return_reference_total": reference.get("total"),
                  "return_critical_qty": reference.get("mean")*profile.return_critical_multiple if number(reference.get("mean")) is not None else None,
                  "iv_score": iv_score, "client_score": client["score"],
                  "product_score": product_score, "supplier_score": supplier["score"],
                  "client_contribution": client["lower"]*profile.client_weight/100,
                  "product_contribution": product_score*profile.product_weight/100 if product_score is not None else 0. if profile.product_weight == 0 else None,
                  "supplier_contribution": supplier["lower"]*profile.supplier_weight/100,
                  "risk_score": lower if math.isclose(lower, upper) else None,
                  "risk_lower": lower, "risk_upper": upper, "known_weight": coverage,
                  "strategy_status": strategy["status"], "strategy_indicator": strategy.get("indicator"), "strategy_reason": strategy["reason"],
                  "strategy_models_matched": strategy["matched_models"], "strategy_models_total": strategy["model_count"],
                  "strategy_source_rows": strategy["source_rows"], "complete": math.isclose(lower, upper)})
    return facts


def build_tu_risk(pqc, aql, voice, iv_cases, scope: DashboardScope, *, grain="Model", department="Hugss",
                  profile=None, registry=(), season="", inspection_only=True):
    """One recomputed score per supplier+grain; no duplicated CC/Model exports."""
    if grain not in GRAINS or department not in DEPARTMENTS:
        raise ValueError("Unsupported TU scope")
    profile = (profile or default_profile(grain)).validate()
    if grain == "Supplier" and profile.product_weight != 0:
        raise ValueError("Supplier grain cannot include product Strategy weight")
    if department != "Hugss" or not scope.includes_supplier("ZX"):
        return pd.DataFrame(), {"department": department, "status": "source_unavailable", "grain": grain}
    p = filter_records(normalize_source(pqc, "pqc"), scope, supplier="ZX", cc_col="cc", model_col="model")
    a = filter_records(normalize_source(aql, "aql"), scope, supplier="ZX", cc_col="cc", model_col="model")
    vgrain = "Model" if grain == "Model" or scope.models else "CC"
    v = voice.loc[voice.get("customer_grain", pd.Series("", index=voice.index)).eq(vgrain)].copy()
    v = filter_records(normalize_source(v, "customer"), scope, supplier="ZX", cc_col="cc", model_col="model", date_col=None)
    iv_available = not iv_cases.empty and {"case_id", "date", "responsibility_stage", "iv_qty"}.issubset(iv_cases.columns)
    i = normalize_source(filter_records(iv_cases, scope, supplier="ZX"), "iv")
    before = i.get("responsibility_stage", pd.Series("", index=i.index)).fillna("").astype(str).str.casefold().str.startswith("before")
    i = i.loc[before].copy()
    # Unassigned IDs are never distributed over all Models/CCs.
    key = {"Model": "model", "CC": "cc", "Supplier": "supplier_code"}[grain]
    # Freeze the peer scope before narrowing to a particular CC/Model. A
    # selected product must not become its own 50-point reference.
    peer_scope = replace(scope, ccs=(), models=())
    peer_p = filter_records(normalize_source(pqc, "pqc"), peer_scope, supplier="ZX", cc_col="cc", model_col="model")
    peer_a = filter_records(normalize_source(aql, "aql"), peer_scope, supplier="ZX", cc_col="cc", model_col="model")
    peer_grain = "Model" if grain in ("Model", "Supplier") else "CC"
    peer_v = voice.loc[voice.get("customer_grain", pd.Series("", index=voice.index)).eq(peer_grain)].copy()
    peer_v = filter_records(normalize_source(peer_v, "customer"), peer_scope, supplier="ZX", cc_col="cc", model_col="model", date_col=None)
    peer_i = normalize_source(filter_records(iv_cases, peer_scope, supplier="ZX"), "iv")
    peer_before = peer_i.get("responsibility_stage", pd.Series("", index=peer_i.index)).fillna("").astype(str).str.casefold().str.startswith("before")
    peer_i = peer_i.loc[peer_before]
    peer_sources = [peer_p, peer_a] if inspection_only else [peer_p, peer_a, peer_v, peer_i]
    peer_population = {str(value) for frame in peer_sources for value in frame[key] if value}
    reference = return_qty_reference(peer_v, peer_population, key)
    sources = [p, a] if inspection_only else [p, a, v, i]
    universe = sorted({str(value) for frame in sources for value in frame[key] if value})
    rows = []
    for value in universe:
        groups = [frame.loc[frame[key].eq(value)] for frame in (p, a, v, i)]
        models = sorted({model for frame in groups for model in frame.model if model})
        ccs = sorted({cc for frame in groups for cc in frame.cc if cc})
        cc = value if grain == "CC" else ccs[0] if len(ccs) == 1 else ""
        strategy = strategic_identity(models, cc, registry, season=season)
        if grain == "Supplier":
            strategy = {"status": "not_applicable", "score": None, "reason": "Strategy belongs to products", "matched_models": 0, "model_count": len(models), "source_rows": []}
        scored = score_entity(*groups, profile, strategy, iv_available=iv_available, rpm_available=scope.period != "Custom", return_reference=reference)
        defect = groups[0].groupby("defect_type")["defect_qty"].sum().sort_values(ascending=False) if "defect_type" in groups[0] and not groups[0].empty else pd.Series(dtype=float)
        rows.append({"entity": value, "grain": grain, "department": department, "supplier_code": "49425", "supplier": "中兴",
                     "cc": cc, "models": " / ".join(models), "top_defect": str(defect.index[0]) if len(defect) else "", **scored})
    result = pd.DataFrame(rows)
    if not result.empty:
        result = result.sort_values(["risk_lower", "entity"], ascending=[False, True]).reset_index(drop=True)
    return result, {"version": RISK_POC_VERSION, "department": department, "supplier_code": "49425", "grain": grain,
                    "entities": len(result), "pqc_rows": len(p), "aql_rows": len(a),
                    "pqc_unmapped_rows": int(p[key].eq("").sum()), "aql_unmapped_rows": int(a[key].eq("").sum()),
                    "iv_available": iv_available, "customer_period": "N0 snapshot; no calendar dates",
                    "client_measure": "Return Qty / PRODUCT QUANTITY", "return_reference": reference,
                    "return_reference_scope": "Same supplier/department/product type/grain and inspection population; individual CC/Model filters excluded",
                    "strategy_source": "Binary Strategy list: listed strategic Model = 1, otherwise = 0", "season": season,
                    "profile": profile.__dict__, "ranking": "known weighted contribution; incomplete bounds do not prove total-risk ordering"}
