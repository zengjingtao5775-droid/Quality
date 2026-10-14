"""Scoring contracts: source grains, configurable weights and unknown inputs."""
from dataclasses import replace
from pathlib import Path
import tempfile
import unittest

import numpy as np
import pandas as pd
from openpyxl import Workbook

from dashboard_scope import DashboardScope
from tu_risk_poc import (RiskProfile, default_profile, threshold_score, weighted_bounds,
                         id_tokens, strategic_identity, load_strategy_workbook,
                         build_tu_risk, aql_facts, pqc_facts)
from quality_insights import build_chart_insight


class RiskPocTests(unittest.TestCase):
    def sources(self):
        pqc=pd.DataFrame({'product_code':['335330','335331'],'model_code':['8655827','8655828'],
                          'date':['2026-06-01','2026-06-02'],'qty_inspected':[55362,1000],
                          'defect_qty':[1136,100],'defect_type':['缝份不匀','线头']})
        aql=pd.DataFrame({'product_code':['335330']*20+['335331'],'model_code':['8655827']*20+['8655828'],
                          'date':['2026-06-01']*21,'record_id':list(range(21)),
                          'result_raw':['PASS']*19+['FAIL','PASS']})
        voice=pd.DataFrame({'customer_grain':['Model','Model','CC','CC'],
                            'product_code':['335330','335331','335330','335331'],
                            'model_code':['8655827','8655828','',''],'returned_now':[162,10,162,10],
                            'sold_now':[180231,10000,180231,10000]})
        iv=pd.DataFrame({'product_code':['335330']*8,'model_code':['8655827']*8,
                         'date':['2026-06-01']*8,'case_id':['a','b','c','d','e','f','f','after'],
                         'responsibility_stage':['Before Sales']*7+['After Sales']})
        return pqc,aql,voice,iv

    def scope(self,**kw):
        return DashboardScope(start='2025-07-02',end='2026-07-01',**kw)

    def test_piecewise_conversion_and_cap(self):
        for value,expected in [(0,0),(2,25),(4,50),(6,75),(8,100),(99,100)]:
            self.assertEqual(threshold_score(value,4,8),expected)
        self.assertIsNone(threshold_score(None,4,8))
        self.assertIsNone(threshold_score(np.inf,4,8))
        self.assertIsNone(threshold_score(-1,4,8))

    def test_nested_weights_are_validated_independently(self):
        for profile in [replace(RiskProfile(),supplier_weight=40),replace(RiskProfile(),iv_weight=30),
                        replace(RiskProfile(),aql_weight=30),replace(RiskProfile(),product_weight=-1,supplier_weight=71),
                        replace(RiskProfile(),rpm_warning=0),replace(RiskProfile(),aql_critical_pct=10)]:
            with self.assertRaises(ValueError):profile.validate()
        self.assertEqual(default_profile('Supplier').product_weight,0)
        self.assertEqual(default_profile('Supplier').supplier_weight,70)

    def test_missing_weight_is_not_zero_or_renormalized(self):
        result=weighted_bounds(((40,60),(None,40)))
        self.assertEqual(result,{'score':None,'lower':24,'upper':64,'coverage':60})
        self.assertEqual(weighted_bounds(((40,100),(None,0)))['score'],40)

    def test_zx_335330_arithmetic_and_strategy_pending(self):
        rows,_=build_tu_risk(*self.sources(),self.scope())
        row=rows.set_index('entity').loc['8655827']
        self.assertAlmostEqual(row.pqc_smoothed,1144/55562)
        self.assertAlmostEqual(row.pqc_score,25.737014506316547)
        self.assertAlmostEqual(row.rpm,162/180231*1e6)
        self.assertAlmostEqual(row.rpm_score,69.5701610178186)
        self.assertEqual(row.iv,6)
        self.assertEqual(row.iv_score,60)
        self.assertAlmostEqual(row.aql_fail_rate,1/20)
        self.assertEqual(row.aql_score,25)
        self.assertAlmostEqual(row.client_score,(row.rpm_score+60)/2)
        self.assertAlmostEqual(row.supplier_score,row.pqc_score*.8+25*.2)
        self.assertEqual(row.strategy_status,'pending')
        self.assertTrue(pd.isna(row.risk_score))
        self.assertEqual(row.known_weight,90)
        self.assertAlmostEqual(row.risk_upper-row.risk_lower,10)
        self.assertAlmostEqual(row.risk_lower,row.client_score*.3+row.supplier_score*.6)

    def test_strategy_exact_yes_no_and_season_supplier_isolation(self):
        registry=[{'models':['8655827'],'ccs':['335330'],'supplier_code':'49425','season':'AW26','strategic':True,'source_row':4}]
        for season,expected in [('AW26','yes'),('SS27','candidate'),('','candidate')]:
            identity=strategic_identity(['8655827'],'335330',registry,season=season)
            self.assertEqual(identity['status'],expected)
        self.assertEqual(strategic_identity(['8655827'],'335330',registry,supplier_code='99999',season='AW26')['status'],'candidate')
        registry[0]['strategic']=False
        self.assertEqual(strategic_identity(['8655827'],'335330',registry,season='AW26')['score'],0)
        self.assertIsNone(strategic_identity(['8655828'],'335330',registry,season='AW26')['score'])

    def test_strategy_conflicting_flags_or_partial_cc_never_complete(self):
        yes={'models':['8655827'],'supplier_code':'49425','season':'AW26','strategic':True}
        no={**yes,'strategic':False}
        self.assertIsNone(strategic_identity(['8655827'],'335330',[yes,no],season='AW26')['score'])
        self.assertIsNone(strategic_identity(['8655827','8655828'],'335330',[yes],season='AW26')['score'])
        self.assertIsNone(strategic_identity(['8655827'],'335330',[{'ccs':['335330'],'supplier_code':'49425','season':'AW26','strategic':True}],season='AW26')['score'])

    def test_confirmed_strategy_changes_only_product_contribution(self):
        registry=[{'models':['8655827'],'supplier_code':'49425','season':'AW26','strategic':False}]
        no,_=build_tu_risk(*self.sources(),self.scope(),registry=registry,season='AW26')
        registry[0]['strategic']=True
        yes,_=build_tu_risk(*self.sources(),self.scope(),registry=registry,season='AW26')
        no=no.set_index('entity').loc['8655827'];yes=yes.set_index('entity').loc['8655827']
        self.assertTrue(no.complete and yes.complete)
        self.assertAlmostEqual(yes.risk_score-no.risk_score,10)
        self.assertEqual(yes.client_score,no.client_score);self.assertEqual(yes.supplier_score,no.supplier_score)

    def test_cc_recomputes_quantities_and_never_adds_model_scores(self):
        p,a,v,i=self.sources()
        p['product_code']='335330';a['product_code']='335330';v['product_code']='335330'
        v.loc[v.customer_grain.eq('CC') & v.returned_now.eq(162),'returned_now']=172
        v.loc[v.customer_grain.eq('CC') & v.sold_now.eq(180231),'sold_now']=190231
        v=v.loc[~(v.customer_grain.eq('CC') & v.returned_now.eq(10))]
        cc,_=build_tu_risk(p,a,v,i,self.scope(),grain='CC')
        self.assertEqual(len(cc),1)
        self.assertEqual(cc.iloc[0].pqc_inspected,56362)
        self.assertEqual(cc.iloc[0].pqc_defects,1236)
        self.assertAlmostEqual(cc.iloc[0].rpm,172/190231*1e6)
        self.assertEqual(cc.iloc[0].aql_records,21)
        models,_=build_tu_risk(p,a,v,i,self.scope(),grain='Model')
        self.assertNotAlmostEqual(cc.iloc[0].risk_lower,models.risk_lower.sum())

    def test_model_filter_derives_qc_identity_before_filtering(self):
        p,a,v,i=self.sources();p['product_label']=p.pop('model_code')+' / BLACK'
        rows,_=build_tu_risk(p,a,v,i,self.scope(models=('8655827',)))
        self.assertEqual(rows.entity.tolist(),['8655827'])
        self.assertEqual(rows.iloc[0].pqc_inspected,55362)
        self.assertEqual(rows.iloc[0].iv,6)

    def test_no_matches_and_empty_departments_stay_empty(self):
        for kw in ({'department':'Knit'},{'department':'Woven'}):
            r,meta=build_tu_risk(*self.sources(),self.scope(),**kw)
            self.assertTrue(r.empty);self.assertEqual(meta['status'],'source_unavailable')
        for scope in [self.scope(ccs=('999999',)),self.scope(suppliers=('FSD',)),self.scope(product_types=('CPT',))]:
            r,_=build_tu_risk(*self.sources(),scope);self.assertTrue(r.empty)

    def test_custom_calendar_has_no_rpm_and_does_not_borrow_factory_total(self):
        r,_=build_tu_risk(*self.sources(),self.scope(period='Custom',models=('8655827',)))
        row=r.iloc[0]
        self.assertTrue(pd.isna(row.rpm))
        self.assertTrue(pd.isna(row.client_score))
        self.assertEqual(row.known_weight,75)
        self.assertAlmostEqual(row.risk_upper-row.risk_lower,25)

    def test_no_iv_source_is_unknown_but_no_matching_cases_is_measured_zero(self):
        p,a,v,i=self.sources()
        rows,_=build_tu_risk(p,a,v,i,self.scope())
        self.assertEqual(rows.set_index('entity').loc['8655828'].iv,0)
        rows,_=build_tu_risk(p,a,v,pd.DataFrame(),self.scope())
        self.assertTrue(rows.iv.isna().all())
        self.assertTrue(rows.client_score.isna().all())

    def test_signed_sales_adjustments_are_summed_not_rate_averaged(self):
        p,a,v,i=self.sources()
        v=v.loc[v.customer_grain.eq('Model')].copy()
        v.loc[1,'sold_now']=-1000
        v.loc[1,'returned_now']=0
        rows,_=build_tu_risk(p,a,v,i,self.scope(models=('8655827','8655828')),grain='Supplier')
        self.assertAlmostEqual(rows.iloc[0].rpm,162/179231*1e6)
        self.assertEqual(rows.iloc[0].strategy_status,'not_applicable')
        self.assertTrue(rows.iloc[0].complete)

    def test_unknown_pqc_denominator_and_aql_status_cannot_impute_zero(self):
        p,a,v,i=self.sources()
        p.loc[0,'qty_inspected']=np.nan
        a.loc[0,'result_raw']='unknown'
        rows,_=build_tu_risk(p,a,v,i,self.scope())
        row=rows.set_index('entity').loc['8655827']
        self.assertTrue(pd.isna(row.pqc_score));self.assertTrue(pd.isna(row.aql_score));self.assertTrue(pd.isna(row.supplier_score))
        self.assertEqual(row.known_weight,30)
        self.assertEqual(row.aql_valid_records,19)

    def test_recheck_pass_is_first_result_failure_and_ids_are_deduplicated(self):
        frame=pd.DataFrame({'record_id':['a','b','b'],'result_raw':['PASS after re-check重验合格','PASS','PASS']})
        facts=aql_facts(frame,RiskProfile())
        self.assertEqual(facts['aql_records'],2)
        self.assertEqual(facts['aql_fail'],1)
        self.assertEqual(facts['aql_fail_rate'],.5)
        frame.loc[2,'result_raw']='FAIL'
        self.assertIsNone(aql_facts(frame,RiskProfile())['aql_score'])

    def test_customer_conflicting_duplicates_are_not_extra_sales(self):
        p,a,v,i=self.sources();extra=v.iloc[[0]].copy();extra['sold_now']=200000
        rows,_=build_tu_risk(p,a,pd.concat([v,extra]),i,self.scope())
        self.assertTrue(pd.isna(rows.set_index('entity').loc['8655827'].rpm))

    def test_unknown_strategy_with_zero_product_weight_can_complete(self):
        profile=replace(RiskProfile(),product_weight=0,supplier_weight=70)
        rows,_=build_tu_risk(*self.sources(),self.scope(),profile=profile)
        self.assertTrue(rows.complete.all())
        self.assertEqual(rows.known_weight.tolist(),[100,100])

    def test_population_union_does_not_require_rpm_to_show_product(self):
        p,a,v,i=self.sources();v=v.loc[v.model_code.ne('8655827') & v.product_code.ne('335330')]
        rows,_=build_tu_risk(p,a,v,i,self.scope())
        self.assertIn('8655827',rows.entity.tolist())
        self.assertTrue(pd.isna(rows.set_index('entity').loc['8655827'].rpm))

    def test_id_lists_reject_embedded_descriptions_and_do_not_infer_relations(self):
        self.assertEqual(id_tokens('8655827 / 8655828',7),frozenset({'8655827','8655828'}))
        self.assertEqual(id_tokens('8655827.0',7),frozenset({'8655827'}))
        self.assertFalse(id_tokens('old model 8655827',7))
        self.assertFalse(id_tokens('8655827-8655828',7))

    def test_workbook_read_only_main_input_and_no_false_department(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'test.xlsx';workbook=Workbook();sheet=workbook.active;sheet.title='Strategic FG homologation'
            sheet.cell(3,9,'Supplier');sheet.cell(3,12,'Conception code');sheet.cell(3,13,'Model code');sheet.cell(3,15,'Season')
            sheet.cell(4,9,'49425 - ZX');sheet.cell(4,12,'335330 / 335331');sheet.cell(4,13,'8655827 / 8655828');sheet.cell(4,15,'AW26')
            workbook.create_sheet('Archive').cell(1,1,'ignored');workbook.save(path);before=path.read_bytes()
            records=load_strategy_workbook(path)
            self.assertEqual(path.read_bytes(),before)
            self.assertEqual(len(records),1)
            self.assertEqual(records[0]['models'],['8655827','8655828'])
            self.assertNotIn('department',records[0]);self.assertNotIn('audit_score',records[0])

    def test_pending_pareto_ai_cannot_claim_full_risk_order(self):
        facts={'evidence':{'kind':'pareto','measure':'risk_score','poc_pending':True,'ranked':[['335330',34.8]],'total':60,'top_share':.58,'shown_share':1}}
        result=build_chart_insight(facts)
        self.assertIn('已知风险贡献',result['finding'])
        self.assertIn('不代表完整风险',result['interpretation'])
        self.assertIn('Strategy',result['action'])


if __name__=='__main__':unittest.main()
